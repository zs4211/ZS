"""Build the Step 2 paper working draft from reproducible diagnostics.

The script clones the selected experiment-analysis DOCX template and preserves
its visual grammar.  It reads only artifacts already produced by
``run_step2_geometry_diagnostics.py`` and does not execute or alter retrieval.
"""

from __future__ import annotations

import json
import math
import subprocess
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
ARTIFACT_DIR = ROOT / "artifacts" / "step2_geometry"
TEMPLATE = Path(
    r"C:\Users\29371\.codex\plugins\cache\openai-curated-remote"
    r"\openai-templates\0.1.1\skills\artifact-template-experiment-analysis"
    r"\assets\reference.docx"
)
OUTPUT = ARTIFACT_DIR / "论文初稿_Step2_数据覆盖与双雷达几何诊断.docx"
GREEN = "137333"
LIGHT_GRAY = "F2F2F2"
GRID_GRAY = "A6A6A6"


def _clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_run_font(run, size: float | None = None, bold: bool | None = None,
                  color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = "Georgia"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _style_document(document: Document) -> None:
    normal = document.styles["normal"]
    normal.font.name = "Georgia"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[name]
        style.font.name = "Georgia"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.color.rgb = RGBColor.from_string(GREEN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def _set_footer(document: Document, short_title: str) -> None:
    for section in document.sections:
        footer = section.footer
        for text_node in footer._element.iter(qn("w:t")):
            if text_node.text and "Report Name" in text_node.text:
                text_node.text = text_node.text.replace("Report Name", short_title)
        for paragraph in footer.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size=8, color=GREEN)


def _body(document: Document, text: str, *, bold_lead: str | None = None,
          italic: bool = False) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        _set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        _set_run_font(rest, italic=italic)
    else:
        run = paragraph.add_run(text)
        _set_run_font(run, italic=italic)


def _bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="normal")
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.16)
    paragraph.paragraph_format.space_after = Pt(3)
    _set_run_font(paragraph.add_run("•  "))
    _set_run_font(paragraph.add_run(text))


def _numbered(document: Document, text: str) -> None:
    counter = getattr(_numbered, "counter", 0) + 1
    _numbered.counter = counter
    paragraph = document.add_paragraph(style="normal")
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.2)
    paragraph.paragraph_format.space_after = Pt(3)
    _set_run_font(paragraph.add_run(f"{counter}.  "), bold=True, color=GREEN)
    _set_run_font(paragraph.add_run(text))


def _heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        _set_run_font(run, color=GREEN, bold=True)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=60, start=80, bottom=60, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRID_GRAY)


def _table(document: Document, headers: list[str], rows: list[list[str]],
           widths: list[float] | None = None, font_size: float = 8.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    if widths:
        total_twips = 9360
        width_sum = sum(widths)
        column_twips = [round(value / width_sum * total_twips) for value in widths]
        column_twips[-1] += total_twips - sum(column_twips)
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.first_child_found_in("w:tblW")
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(total_twips))
        tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:type"), "dxa")
        tbl_ind.set(qn("w:w"), "80")
        for grid_column, width_twips in zip(table._tbl.tblGrid.gridCol_lst, column_twips):
            grid_column.set(qn("w:w"), str(width_twips))
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for col, value in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = str(value)
        _set_cell_shading(cell, LIGHT_GRAY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margins(cell)
        if widths:
            cell.width = Twips(column_twips[col])
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size=font_size, bold=True, color=GREEN)
    for values in rows:
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cell = cells[col]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)
            if widths:
                cell.width = Twips(column_twips[col])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    _set_run_font(run, size=font_size, bold=(col == 0))
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(7)
    _set_run_font(paragraph.add_run(text), size=8.5, italic=True, color="555555")


def _figure(document: Document, path: Path, caption: str, width: float = 5.8) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    inline = paragraph.add_run().add_picture(str(path), width=Inches(width))
    inline._inline.docPr.set("descr", caption)
    inline._inline.docPr.set("title", caption.split("。", 1)[0])
    _caption(document, caption)


def _page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _fmt(value, digits=3) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if not math.isfinite(number):
        return "N/A"
    return f"{number:.{digits}f}"


def _pct(value, digits=2) -> str:
    return f"{100.0 * float(value):.{digits}f}%"


def _git_head() -> str:
    return subprocess.check_output(
        ["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, text=True
    ).strip()


def _cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(28)
    _set_run_font(paragraph.add_run("SCIENTIFIC DIAGNOSTIC REPORT"), size=13, color=GREEN)
    divider = document.add_paragraph("────────")
    divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(divider.runs[0], size=11, color=GREEN)
    for _ in range(3):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True
    _set_run_font(title.add_run("结构性缺测下双多普勒三维风场反演"), size=25, bold=True, color=GREEN)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(subtitle.add_run("数据覆盖、Mask 来源与几何可观测性诊断"), size=21, bold=True, color=GREEN)
    label = document.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(label.add_run("论文初稿 · Step 2"), size=14, color=GREEN)
    for _ in range(10):
        document.add_paragraph()
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(author.add_run("ZS 项目 / Codex 实验记录整理"), size=10.5)
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(date.add_run("2026 年 9 月 4 日"), size=10.5)
    _page_break(document)


def build() -> Path:
    summary = json.loads((ARTIFACT_DIR / "step2_summary.json").read_text(encoding="utf-8"))
    height = pd.read_csv(ARTIFACT_DIR / "coverage_by_height.csv")
    bca = pd.read_csv(ARTIFACT_DIR / "bca_bins_by_height.csv")
    ranges = pd.read_csv(ARTIFACT_DIR / "range_geometry_summary.csv")
    features = pd.read_csv(ARTIFACT_DIR / "feature_neighbourhood_diagnostics.csv")
    repeated = pd.read_csv(ARTIFACT_DIR / "repeated_sweep_diagnostics.csv")
    document = Document(str(TEMPLATE))
    _clear_body(document)
    _style_document(document)
    _set_footer(document, "Step 2 Geometry Diagnostic")
    document.core_properties.title = "数据覆盖、Mask 来源与双雷达几何可观测性诊断"
    document.core_properties.subject = "雷达三维风场张量反演论文初稿，Step 2"
    document.core_properties.author = "ZS / Codex"

    _cover(document)
    _heading(document, "文档控制")
    _table(document, ["字段", "记录"], [
        ["版本", "0.2 — Step 2 诊断初稿"],
        ["状态", "工作草稿；等待 Step 2 人工验收"],
        ["冻结基线", "step0-baseline (082209f)"],
        ["分析代码起点", f"Git HEAD {_git_head()}（生成时工作区含本轮未提交结果）"],
        ["案例", "Z9539 2026-05-11 07:09:39 UTC / Z9532 07:09:29 UTC"],
        ["网格", "201 × 201 × 20；x,y=[−150,150] km；z=[0.5,15] km"],
        ["本轮边界", "仅覆盖、mask provenance、双雷达几何；未改科学算法"],
    ], widths=[1.5, 5.0], font_size=9)

    _heading(document, "摘要")
    grid = summary["grid_coverage"]
    _body(document, (
        "本研究面向结构性缺测与不良双多普勒几何下的三维风场反演。本轮以冻结的 "
        "Z9539/Z9532 个例为对象，独立诊断原始门数据、现有 QC 后数据、网格化覆盖、"
        "双雷达交叠区、对称交叉波束角（BCA）、水平观测矩阵奇异值以及 2×3 直接观测"
        "子空间的零空间方向。结果显示，每部雷达在全域网格上的速度覆盖分别为 "
        f"{_pct(grid['radar1_fraction'])} 与 {_pct(grid['radar2_fraction'])}，但双雷达交叠仅 "
        f"{_pct(grid['dual_before_bca_fraction'])}。现有 BCA≥20° 门槛进一步剔除交叠区的 "
        f"{_pct(grid['rejected_by_bca_fraction_of_dual'])}，最终有效比例为 "
        f"{_pct(grid['bca_ge_20_fraction'], 3)}。因此 0.86% 的主要瓶颈是覆盖重叠不足，"
        "BCA 门槛是交叠区内显著的二次损失。两处当前反演涡旋候选均处于比可解析 M 产品"
        "特征更好的双雷达几何区，但这只说明其直接观测条件更好，不证明其为真实中气旋。"
    ))
    _body(document, "关键词：双多普勒雷达；三维风场反演；结构性缺测；交叉波束角；可观测性；mask provenance")

    _heading(document, "1 研究问题与 Step 2 范围")
    _body(document, (
        "第一篇论文的总问题是：时空低秩结构与质量连续约束能否在结构性缺测和不良双多普勒"
        "几何条件下提高三维风反演稳定性，同时保持局地涡旋结构。Step 2 不评价 Tucker 或"
        "物理正则化效果，只建立之后所有实验必须共享的数据覆盖与直接雷达可观测性基线。"
    ))
    _bullet(document, "未修改 PyDDA 求解器、滤波参数、速度/反射率 QC 规则、中气旋阈值或 BCA=20° 门槛。")
    _bullet(document, "026/027 重复仰角均按现状保留；只审计，不删除、不合并。")
    _bullet(document, "业务 M 产品仅作为位置与旋转信号的辅助参考，不作为三维风场真值。")

    _heading(document, "2 数学约定与独立几何指标")
    _body(document, "在 Step 1 验证的 ENU 约定下，x 向东、y 向北、z 向上；方位角自正北顺时针；仰角向上为正；径向速度正值表示远离雷达。")
    _body(document, "对雷达 r，单位波束向量 hᵣ=[cos(eᵣ)sin(aᵣ), cos(eᵣ)cos(aᵣ), sin(eᵣ)]ᵀ。双雷达单格点观测矩阵 Hᵢ=[h₁ᵀ;h₂ᵀ]∈R²ˣ³，故 rank(Hᵢ)≤2。本文只报告两个非零奇异值与零空间单位方向 nᵢ，并以 |nᵢ·e_z| 描述不可观方向和垂直方向的一致性。普通三维 condition number 在这里必然为无穷，不被用作质量指标。")
    _body(document, "水平矩阵 A_h 由两部雷达波束的东、北分量构成。报告 σmax(A_h)、σmin(A_h)、σmin/σmax 与 κ₂(A_h)。对称 BCA 定义为两条视线无向夹角，因此原始夹角 170° 映射为 10°。单雷达或无雷达格点的全部双雷达指标保持 N/A。")

    _heading(document, "3 数据与方法")
    _table(document, ["环节", "实际输入/处理", "本轮记录"], [
        ["极坐标速度", "Vc；10 个 sweep/雷达", "原始存在、|Vr|≤100 m s⁻¹ 基础范围检查、现有 QC 后存在"],
        ["反射率", "Zc；10 个 sweep/雷达", "原始与现有 QC 后存在数量"],
        ["网格化", "现有 PyART 路径，201×201×20", "每部雷达有效 mask 与共同有效 mask"],
        ["几何", "独立模块，matrix-free", "BCA、水平 SVD、2×3 H 非零 SVD、零空间"],
        ["最终 mask", "现有 BCA≥20° + 下游 mask", "逐格与冻结基线对账"],
    ], widths=[1.0, 2.0, 3.5], font_size=8.4)

    _heading(document, "4 总体覆盖与 mask 归因结果")
    raw = summary["polar_gate_provenance"]
    _table(document, ["雷达", "原始 Vc 存在", "现有 QC 后 Vc", "范围 QC 失败", "原始/现有 QC 后 Zc"], [
        [station,
         f"{raw['raw_velocity'][station]['present_gates']:,} ({_pct(raw['raw_velocity'][station]['present_fraction'])})",
         f"{raw['after_unchanged_qc'][station]['present_gates']:,} ({_pct(raw['after_unchanged_qc'][station]['present_fraction'])})",
         f"{raw['raw_velocity'][station]['range_qc_fail_gates']:,}",
         f"{raw['reflectivity'][station]['raw_present_gates']:,} / {raw['reflectivity'][station]['qc_present_gates']:,}"]
        for station in ("Z9539", "Z9532")
    ], widths=[0.8, 1.5, 1.5, 1.1, 1.6], font_size=8.2)
    classes = summary["observation_classes"]
    _table(document, ["类别", "格点数", "占全域"], [
        ["双雷达好几何（BCA≥20°）", f"{classes['dual_good_geometry_bca_ge_20']:,}", _pct(classes['dual_good_geometry_bca_ge_20']/summary['case']['total_gridpoints'], 3)],
        ["双雷达差几何（BCA<20°）", f"{classes['dual_poor_geometry_bca_lt_20']:,}", _pct(classes['dual_poor_geometry_bca_lt_20']/summary['case']['total_gridpoints'], 3)],
        ["仅单雷达覆盖", f"{classes['single_radar']:,}", _pct(classes['single_radar']/summary['case']['total_gridpoints'])],
        ["完全无速度观测", f"{classes['no_observation']:,}", _pct(classes['no_observation']/summary['case']['total_gridpoints'])],
    ], widths=[3.5, 1.2, 1.2], font_size=9)
    _body(document, f"独立诊断的最终 mask 与冻结基线逐格不一致数为 {grid['diagnostic_vs_baseline_disagreement_count']}；BCA 之后未发现额外 downstream 排除格点。")
    _figure(document, ARTIFACT_DIR / "coverage_classes_4km.png", "图 1  z=4.32 km 的四类观测区域。黄色为单雷达覆盖，红/绿分别为差/好双雷达几何。", width=5.4)

    _heading(document, "5 按高度、BCA 与距离分层")
    nonzero_height = height[height["radar1_fraction"].gt(0) | height["radar2_fraction"].gt(0)]
    _table(document, ["高度/km", "Z9539", "Z9532", "双雷达/BCA前", "BCA≥20°", "门槛剔除/双雷达"], [
        [_fmt(row.height_m/1000, 2), _pct(row.radar1_fraction), _pct(row.radar2_fraction),
         _pct(row.dual_before_bca_fraction), _pct(row.bca_ge_20_fraction),
         "N/A" if pd.isna(row.bca_rejected_fraction_of_dual) else _pct(row.bca_rejected_fraction_of_dual)]
        for row in nonzero_height.itertuples()
    ], widths=[0.75, 0.85, 0.85, 1.3, 1.05, 1.35], font_size=7.6)
    _figure(document, ARTIFACT_DIR / "coverage_by_height.png", "图 2  各高度层单雷达、双雷达与现有 BCA 门槛后的覆盖比例。", width=5.9)

    bca_total = bca.groupby("bca_bin_deg", sort=False)["gridpoint_count"].sum()
    _table(document, ["对称 BCA", "格点数", "占双雷达区", "占全域"], [
        [str(bin_name), f"{int(count):,}", _pct(count/bca_total.sum()), _pct(count/summary['case']['total_gridpoints'], 3)]
        for bin_name, count in bca_total.items()
    ], widths=[1.4, 1.2, 1.4, 1.2], font_size=8.6)
    _table(document, ["两雷达最大距离/km", "域格点", "双雷达", "双雷达率", "BCA≥20°/双雷达", "中位 BCA", "中位 σmin/σmax"], [
        [str(row.max_radar_range_bin_km), f"{int(row.domain_count):,}", f"{int(row.dual_count):,}",
         _pct(row.dual_fraction_in_range_bin) if math.isfinite(row.dual_fraction_in_range_bin) else "N/A",
         _pct(row.bca_ge_20_fraction_of_dual) if math.isfinite(row.bca_ge_20_fraction_of_dual) else "N/A",
         _fmt(row.median_bca_deg, 2), _fmt(row.median_horizontal_inverse_condition, 3)]
        for row in ranges.itertuples()
    ], widths=[1.2, 0.9, 0.8, 1.0, 1.25, 0.9, 1.15], font_size=7.5)
    _body(document, "双雷达交叠仅出现在“两雷达最大距离”为 75–125 km 的分层中；125 km 以上虽然仍可能有单雷达网格值，但本案例中没有共同有效速度格点。")

    _heading(document, "6 业务 M 产品位置诊断")
    m_features = features[features["feature_type"].eq("business_M")]
    _table(document, ["位置", "x/y/z km", "R1/R2 km", "有效 R1/R2", "邻域双雷达", "邻域中位 BCA", "进入现有有效区"], [
        [row.feature_id,
         f"{_fmt(row.x_km,1)}/{_fmt(row.y_km,1)}/{_fmt(row.z_km,2)}",
         f"{_fmt(row.radar1_range_km,1)}/{_fmt(row.radar2_range_km,1)}",
         f"{'是' if row.radar1_valid else '否'}/{'是' if row.radar2_valid else '否'}",
         _pct(row.neighbourhood_dual_fraction), _fmt(row.neighbourhood_median_bca_deg,2),
         "是" if row.current_pydda_effective_dual else "否"]
        for row in m_features.itertuples()
    ], widths=[0.85, 1.25, 1.05, 1.0, 1.0, 1.05, 1.15], font_size=7.2)
    m_reason_cn = {
        "Z9539_M1": "两部雷达目标邻近极坐标窗口均无有效速度，网格化后无观测",
        "Z9539_M2": "Z9532 目标距离超过当前速度记录量程",
        "Z9532_M1": "目标对两部雷达均超过当前速度记录量程",
        "Z9532_M2": "Z9539 量程内邻近窗口无有效 QC 后速度；邻域双雷达 BCA 仍低于 20°",
        "Z9532_M3": "Z9532 目标距离超过当前速度记录量程",
    }
    _table(document, ["位置", "具体失效原因（现有数据可支持的最细归因）"], [
        [row.feature_id, m_reason_cn[row.feature_id]] for row in m_features.itertuples()
    ], widths=[1.0, 5.6], font_size=8.3)
    _body(document, (
        "结论不是“所有 M 位置都由同一原因缺失第二部雷达”。西侧 Z9539_M2 与 Z9532_M3 的 "
        "Z9532 目标距离约 124–125 km，超过当前速度资料记录量程；远西北 Z9532_M1 对两部雷达"
        "均超出记录量程。中央高层 Z9539_M1 位于量程内，但两部雷达目标邻近极坐标窗口均无"
        "有效速度。Z9532_M2 位于量程内，目标格点缺 Z9539，且其 5 km 邻域虽有 18.18% 双雷达"
        "覆盖，中位 BCA 仅 15.30°，全部被现有 20° 门槛排除。"
    ))
    _body(document, "由于 PyART Grid 未保存每个格点的源 gate ID、ROI 候选及插值权重，上述“邻近极坐标有/无有效样本”不能在所有点位上唯一还原成某一条门级 QC 或插值规则。不能对已丢失的 provenance 作猜测。")

    _heading(document, "7 当前反演涡旋候选位置诊断")
    candidates = features[features["feature_type"].eq("retrieval_vortex_candidate")]
    _table(document, ["候选", "x/y/z km", "ζ /s", "dBZ", "BCA", "σmin/σmax", "κ₂(Ah)", "3D σmin/σmax", "|n·ez|", "有效"], [
        [row.feature_id,
         f"{_fmt(row.x_km,1)}/{_fmt(row.y_km,1)}/{_fmt(row.z_km,2)}",
         _fmt(row._8, 6),
         _fmt(row.reflectivity_dbz,1), _fmt(row.symmetric_bca_deg,2),
         _fmt(row.horizontal_inverse_condition,3), _fmt(row.horizontal_condition_number,3),
         f"{_fmt(row.direct_3d_sigma_min,3)}/{_fmt(row.direct_3d_sigma_max,3)}",
         _fmt(row.null_vertical_alignment,3), "是" if row.current_pydda_effective_dual else "否"]
        for row in candidates.itertuples(index=False)
    ], widths=[0.65, 1.05, 0.75, 0.55, 0.65, 0.85, 0.75, 1.05, 0.65, 0.5], font_size=6.7)
    _body(document, "C1 与 C2 的 BCA 分别为 36.33° 与 47.11°，水平 κ₂ 分别为 3.05 与 2.29；其 5 km 邻域中位 BCA 分别为 34.95° 与 46.11°。它们显著优于可解析 M 位置的无双雷达或约 15° 差几何，但 |n·e_z|≈0.99 表明两部近水平波束的直接不可观方向仍几乎垂直，w 不能由单格点双雷达直接辨识。")
    _body(document, "两候选峰值涡度约 0.00184 s⁻¹ 与 0.00165 s⁻¹，且反射率约 19.4/22.4 dBZ；本轮未降低 0.005 s⁻¹ 阈值，也不把它们认定为业务中气旋。")

    _heading(document, "8 重复 sweep 审计")
    repeated_only = repeated[repeated["is_repeated_elevation"]]
    _table(document, ["雷达", "时刻 UTC", "仰角", "目录", "射线×距离门", "有效门比例"], [
        [row.station, row.scan_time_utc.replace("+00:00", "Z"), _fmt(row.elevation_deg,1),
         f"{int(row.product_directory):03d}", f"{int(row.rays)}×{int(row.gates_per_ray)}", _pct(row.valid_gate_fraction)]
        for row in repeated_only.itertuples()
    ], widths=[0.75, 1.65, 0.65, 0.65, 1.15, 1.0], font_size=7.5)
    _body(document, "两部雷达均有 3 组重复实际仰角（0.5°、1.5°、2.4°），每组保留 026 与 027 两个 sweep，共 6 条重复 sweep 记录/雷达。所有记录时间均为各自体扫时刻；026 的有效门比例系统性高于同仰角 027。")
    _body(document, "当前转换把每个 sweep 作为独立 ray 集合传给 PyART。归一化插值通常不会简单把速度幅值相加两次，但重复采样会改变邻域候选门数量及有效权重，并可能扩展有效 mask 或改变插值结果。本轮未合并它们，因而这里只能标记为后续敏感性实验风险。")

    _heading(document, "9 四个科学问题的直接回答")
    _numbered(document, "0.86% 主要由双雷达覆盖不足主导。BCA 前双雷达覆盖只有 1.514%；98.486% 全域格点本来就不是双雷达共同有效。20° 门槛又移除 5,306 个格点，即双雷达区的 43.37%，是显著但次级的损失。")
    _numbered(document, "M 产品位置缺失第二雷达的原因具有位置依赖性：有的超过 Z9532 或两部雷达当前速度记录量程；有的在高层处于极坐标有效速度空洞；有的存在少量双雷达邻域覆盖但 BCA≈15°，被现有门槛排除。门到格点映射丢失使个别量程内点位无法再唯一追溯到具体 QC/ROI/插值规则。")
    _numbered(document, "两个反演涡旋候选确实位于更好的双雷达水平几何区域（BCA 36.3°/47.1°，κ₂ 3.05/2.29），而可解析 M 位置要么没有双雷达覆盖，要么邻域 BCA 约 15.3°。这只能解释“哪里更容易被当前双雷达流程形成解”，不能证明候选是真实涡旋。")
    _numbered(document, "区域分类已逐格输出：双雷达好几何 6,928；双雷达差几何 5,306；单雷达 144,130；完全无观测 651,656。NetCDF 保存每类 mask 及全部几何变量，CSV/图给出高度、BCA 与距离分层。")

    _heading(document, "10 数据质量、局限与下一步决策")
    _bullet(document, "当前真实案例只有一个时次；结果不可泛化为所有雷达对或天气过程。")
    _bullet(document, "M 产品不是三维真值，反射率也不能证明反演风正确。")
    _bullet(document, "两部雷达单格点最多直接约束二维子空间；后续质量连续、时间约束和 Tucker 先验用于限制剩余方向，但不能创造不存在的信息。")
    _bullet(document, "距离分层以两雷达距离最大值定义；高度层表中的百分比均以该层 40,401 个水平格点为分母。")
    _bullet(document, "未来应在读取/QC/网格化阶段持久化 gate-level reason code、源 gate ID、ROI、权重和 contributor count，以恢复完整 provenance；该建议不在 Step 2 内实施。")
    _body(document, "Step 2 判定：通过。独立诊断与冻结 mask 一致、四个科学问题得到数据支持、24 项 unittest 与 compileall 均通过。按既定顺序，等待人工确认后才可进入 Step 3 的 Dρ/Dρ*；本文档不宣称 Step 3 已开始。")

    _heading(document, "附录 A：复现实验")
    _table(document, ["目的", "命令", "结果"], [
        ["重建 Step 2 真实案例", "python tools/run_step2_geometry_diagnostics.py", "成功；输出 NetCDF/CSV/JSON/PNG"],
        ["几何/mask 定向测试", "python -m unittest tests.test_geometry_diagnostics -v", "6/6 通过"],
        ["完整测试", "python -m unittest discover -s tests -v", "24/24 通过"],
        ["语法编译", "python -m compileall -q src tests tools", "通过"],
    ], widths=[1.3, 3.8, 1.4], font_size=8.2)
    _body(document, "主数据产品：artifacts/step2_geometry/geometry_diagnostics_20260511070939Z_20260511070929Z.nc。摘要、分层表、点位表与重复 sweep 表位于同目录。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    print(build())
